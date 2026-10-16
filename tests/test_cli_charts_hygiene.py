from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
import pytest

from lbo_lab import charts, sensitivity
from lbo_lab.assumptions import Assumptions
from lbo_lab.cli import main
from lbo_lab.engine import run_model

ROOT = Path(__file__).resolve().parents[1]
DASH = re.compile("[" + chr(0x2013) + chr(0x2014) + "]")  # en and em dashes are banned


def test_cli_run_prints_sections(capsys):
    assert main(["run"]) == 0
    out = capsys.readouterr().out
    for text in ("SOURCES AND USES", "RETURNS", "VALUE CREATION BRIDGE", "Sponsor IRR", "Total debt / EBITDA"):
        assert text in out


@pytest.mark.parametrize("scenario", ["base", "management", "downside"])
def test_cli_run_json_matches_engine(capsys, scenario):
    assert main(["run", "--scenario", scenario, "--json"]) == 0
    data = json.loads(capsys.readouterr().out)
    expected = run_model(Assumptions(), scenario).returns
    assert data["sponsor_irr"] == pytest.approx(expected.irr)
    assert data["sponsor_moic"] == pytest.approx(expected.moic)


def test_cli_exit_year_and_pik_flags(capsys):
    assert main(["run", "--json", "--exit-year", "7", "--pik"]) == 0
    data = json.loads(capsys.readouterr().out)
    assert data["exit_year"] == 7


def test_cli_average_interest(capsys):
    assert main(["run", "--interest-method", "average"]) == 0
    assert "average balances" in capsys.readouterr().out


def test_cli_rejects_unknown_scenario(capsys):
    assert main(["run", "--scenario", "moonshot"]) == 2
    assert "unknown scenario" in capsys.readouterr().err


def test_cli_rejects_bad_exit_year(capsys):
    assert main(["run", "--exit-year", "12"]) == 2


def test_cli_scenarios(capsys):
    assert main(["scenarios"]) == 0
    assert "downside" in capsys.readouterr().out


def test_cli_sens_prints_all_grids(capsys):
    assert main(["sens"]) == 0
    out = capsys.readouterr().out
    assert out.count("sponsor IRR") == 3 and out.count("sponsor MOIC") == 3
    assert "Leverage vs exit year" in out


def test_cli_build_without_recalc(tmp_path, capsys):
    out = tmp_path / "sub" / "m.xlsx"
    assert main(["build", "--scenario", "downside", "--out", str(out), "--no-recalc"]) == 0
    assert out.exists()
    assert openpyxl.load_workbook(out)["Inputs"]["D5"].value == "Downside"


def test_cli_charts_writes_pngs(tmp_path, capsys):
    assert main(["charts", "--out-dir", str(tmp_path)]) == 0
    pngs = sorted(p.name for p in tmp_path.glob("*.png"))
    assert pngs == ["debt_paydown.png", "leverage_coverage.png", "sensitivity_irr.png",
                    "value_creation_bridge.png"]
    assert all(p.stat().st_size > 10_000 for p in tmp_path.glob("*.png"))


def test_chart_functions_accept_every_scenario(tmp_path):
    for key in ("base", "management", "downside"):
        a = Assumptions().with_scenario(key)
        paths = charts.make_all(run_model(a), sensitivity.entry_vs_exit(a), tmp_path / key)
        assert len(paths) == 4


# ------------------------------------------------------------------ repository hygiene
def _text_files():
    skip = {".git", "__pycache__", ".pytest_cache", ".ruff_cache", "lbo_lab.egg-info", ".venv"}
    for p in ROOT.rglob("*"):
        if p.is_file() and not (set(p.relative_to(ROOT).parts) & skip) and p.suffix in {
                ".py", ".md", ".toml", ".yml", ".yaml", ".txt", ""}:
            yield p


def test_no_em_or_en_dashes_anywhere():
    bad = [str(p.relative_to(ROOT)) for p in _text_files()
           if p.name != "LICENSE" and DASH.search(p.read_text(encoding="utf-8", errors="ignore"))]
    assert bad == []


def test_no_dashes_in_workbook_text():
    path = ROOT / "examples" / "halcyon_lbo.xlsx"
    if not path.exists():
        pytest.skip("example workbook not generated")
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    assert not DASH.search(c.value), (ws.title, c.coordinate)


def test_no_todo_placeholders_or_contact_details():
    for p in _text_files():
        if p.suffix == ".py" and p.parent.name == "tests":
            continue
        text = p.read_text(encoding="utf-8", errors="ignore")
        assert "TODO" not in text and "FIXME" not in text, p
        assert not re.search(r"[\w.]+@[\w.]+\.\w+", text), p


def test_license_and_gitignore_present():
    assert "Copyright (c) 2026 Adham Karim" in (ROOT / "LICENSE").read_text()
    assert "__pycache__" in (ROOT / ".gitignore").read_text()
